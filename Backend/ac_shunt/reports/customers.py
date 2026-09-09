"""Customer directory workbook parsing and atomic replacement."""
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from django.db import transaction
from .models import Customer, CustomerDirectory

HEADERS = ('Activity', 'Sub Custodian', 'Lab Name', 'Lab Address', 'SYSCOM')
FIELDS = ('activity', 'sub_custodian', 'lab_name', 'address', 'syscom')
MAX_ROWS = 50000


def parse_customers(upload):
    if not str(upload.name).lower().endswith('.xlsx'):
        raise ValueError('Choose an .xlsx Lab Address List workbook.')
    if upload.size > 10 * 1024 * 1024:
        raise ValueError('Workbook must be smaller than 10 MB.')
    try:
        with ZipFile(upload) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Workbook expands beyond the 50 MB limit.')
        upload.seek(0)
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError('Cannot read this Excel workbook. Choose a valid .xlsx file.') from exc
    rows, seen = [], set()
    duplicates = 0
    try:
        found = False
        for sheet in workbook:
            if sheet.max_column and sheet.max_column > 100:
                raise ValueError('Workbook sheets must contain no more than 100 columns.')
            indices = None
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS + 20:
                    raise ValueError('Workbook exceeds the 50,000-row limit.')
                values = [str(cell.value or '').strip() for cell in cells]
                if indices is None:
                    normalized = [value.casefold() for value in values]
                    if all(header.casefold() in normalized for header in HEADERS):
                        indices = [normalized.index(header.casefold()) for header in HEADERS]
                        found = True
                    elif number >= 20:
                        break
                    continue
                record = []
                for index in indices:
                    cell = cells[index] if index < len(cells) else None
                    if cell is not None and cell.data_type == 'f':
                        raise ValueError(f'{sheet.title}, row {number}: replace formulas with values before importing.')
                    value = '' if cell is None or cell.value is None else str(cell.value).strip()
                    # Preserve identifiers stored as numbers with an explicit zero mask.
                    if cell is not None and isinstance(cell.value, (int, float)) and cell.value == int(cell.value):
                        value = str(int(cell.value))
                        if cell.number_format and set(cell.number_format) == {'0'}:
                            value = value.zfill(len(cell.number_format))
                    if len(value) > 255:
                        raise ValueError(f'{sheet.title}, row {number}: fields must be at most 255 characters.')
                    record.append(value)
                if not any(record):
                    continue
                if not any(record[:3]):
                    raise ValueError(f'{sheet.title}, row {number}: customer name or activity/sub-custodian code is required.')
                key = tuple(record)
                if key in seen:
                    duplicates += 1
                    continue
                seen.add(key)
                rows.append(dict(zip(FIELDS, record)))
                if len(rows) > MAX_ROWS:
                    raise ValueError('Workbook exceeds the 50,000-customer limit.')
        if not found:
            raise ValueError('Required headers: Activity, Sub Custodian, Lab Name, Lab Address, SYSCOM.')
        if not rows:
            raise ValueError('No customers found. The existing directory was kept.')
        return rows, duplicates
    finally:
        workbook.close()


def replace_customers(rows, filename):
    # Directory writes and its metadata commit together on the reports database.
    with transaction.atomic(using='reports'):
        directory = CustomerDirectory.objects.using('reports').select_for_update().get(pk=1)
        Customer.objects.using('reports').all().delete()
        Customer.objects.using('reports').bulk_create([Customer(**row) for row in rows], batch_size=500)
        directory.source_name = Path(filename.replace('\\', '/')).name[:255]
        directory.row_count = len(rows)
        directory.save(using='reports')
    return directory
